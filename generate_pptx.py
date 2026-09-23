#!/usr/bin/env python3
"""Generate Brahmora / AEGIS / PRE presentation as PowerPoint."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Brand colours
BRAND = RGBColor(0x63, 0x66, 0xF1)
BRAND_LIGHT = RGBColor(0x81, 0x8C, 0xF8)
CYAN = RGBColor(0x22, 0xD3, 0xEE)
EMERALD = RGBColor(0x34, 0xD3, 0x99)
AMBER = RGBColor(0xFB, 0xBF, 0x24)
ROSE = RGBColor(0xFB, 0x71, 0x85)
VIOLET = RGBColor(0xA7, 0x8B, 0xFA)

BG = RGBColor(0x05, 0x08, 0x0E)
SURFACE = RGBColor(0x0F, 0x16, 0x29)
TEXT1 = RGBColor(0xF0, 0xF4, 0xFC)
TEXT2 = RGBColor(0x94, 0xA3, 0xB8)
TEXT3 = RGBColor(0x64, 0x74, 0x8B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

W = prs.slide_width
H = prs.slide_height


def set_bg(slide, color=BG):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_textbox(slide, left, top, width, height, text, font_size=18,
                color=TEXT1, bold=False, alignment=PP_ALIGN.LEFT,
                font_name="Calibri"):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return txBox


def add_para(text_frame, text, font_size=14, color=TEXT2, bold=False,
             alignment=PP_ALIGN.LEFT, space_before=Pt(6), font_name="Calibri"):
    p = text_frame.add_paragraph()
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    if space_before:
        p.space_before = space_before
    return p


def add_label(slide, left, top, text, color=BRAND_LIGHT):
    add_textbox(slide, left, top, Inches(5), Inches(0.4), text,
                font_size=11, color=color, bold=True)


def add_rounded_rect(slide, left, top, width, height, fill_color=SURFACE,
                     border_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                   left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape


def add_accent_bar(slide, left, top, width, color=BRAND):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Pt(3))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


# ════════════════════════════════════════════════════════════
# SLIDE 1 — TITLE
# ════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
set_bg(slide)

add_textbox(slide, Inches(1.5), Inches(0.8), Inches(10), Inches(0.4),
            "BRAHMORA TECHNOLOGIES", font_size=13, color=BRAND_LIGHT, bold=True)

add_textbox(slide, Inches(1.5), Inches(1.5), Inches(10.3), Inches(2),
            "The traditional way of operating\ncloud platforms no longer scales.",
            font_size=40, color=TEXT1, bold=True, alignment=PP_ALIGN.LEFT)

tb = add_textbox(slide, Inches(1.5), Inches(3.7), Inches(8), Inches(1.5),
                 "Platform teams manage more infrastructure, more risk, more cost, and more governance than ever before. But they still operate with disconnected tools and manual processes.",
                 font_size=16, color=TEXT2)

# Statement box
add_rounded_rect(slide, Inches(1.5), Inches(5.2), Inches(10.3), Inches(1.6), SURFACE, BRAND)
add_accent_bar(slide, Inches(1.5), Inches(5.2), Inches(10.3), BRAND)
add_textbox(slide, Inches(1.8), Inches(5.4), Inches(9.7), Inches(0.6),
            "AEGIS brings reliability, governance, operations, and cost into one operational system.",
            font_size=18, color=TEXT1, bold=True)
add_textbox(slide, Inches(1.8), Inches(5.95), Inches(9.7), Inches(0.5),
            "It works independently as a platform reliability control plane and integrates with existing tools to provide a unified operating layer.",
            font_size=14, color=TEXT2)


# ════════════════════════════════════════════════════════════
# SLIDE 2 — OPERATING MODEL SHIFT
# ════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)

add_label(slide, Inches(1), Inches(0.6), "THE OPERATING MODEL SHIFT")
add_textbox(slide, Inches(1), Inches(1.0), Inches(11), Inches(1.2),
            "Cloud platforms became complex faster\nthan operations evolved",
            font_size=32, color=TEXT1, bold=True)
add_textbox(slide, Inches(1), Inches(2.3), Inches(8), Inches(0.5),
            "Every wave of infrastructure maturity produced a discipline to manage it.",
            font_size=16, color=TEXT2)

# Evolution boxes
evo_data = [("2008", "DevOps", TEXT3), ("2016", "SRE", TEXT3),
            ("2022", "Platform Eng", TEXT3), ("Now", "PRE", BRAND_LIGHT)]
x_start = Inches(1)
for i, (year, name, color) in enumerate(evo_data):
    x = x_start + Inches(i * 3)
    is_active = name == "PRE"
    fill = RGBColor(0x15, 0x1D, 0x30) if not is_active else RGBColor(0x1A, 0x1E, 0x3A)
    border = None if not is_active else BRAND
    add_rounded_rect(slide, x, Inches(3.2), Inches(2.3), Inches(1.0), fill, border)
    add_textbox(slide, x, Inches(3.25), Inches(2.3), Inches(0.35),
                year, font_size=11, color=BRAND_LIGHT if is_active else TEXT3,
                bold=True, alignment=PP_ALIGN.CENTER)
    add_textbox(slide, x, Inches(3.55), Inches(2.3), Inches(0.5),
                name, font_size=18, color=TEXT1 if is_active else TEXT3,
                bold=True, alignment=PP_ALIGN.CENTER)
    if i < 3:
        add_textbox(slide, x + Inches(2.3), Inches(3.45), Inches(0.7), Inches(0.5),
                    "\u2192", font_size=20, color=TEXT3, alignment=PP_ALIGN.CENTER)

# Statement
add_rounded_rect(slide, Inches(1), Inches(4.8), Inches(11.3), Inches(1.8), SURFACE)
add_accent_bar(slide, Inches(1), Inches(4.8), Inches(11.3), BRAND)
add_textbox(slide, Inches(1.4), Inches(5.0), Inches(10.5), Inches(0.7),
            "Platforms must now operate like products \u2014 with reliability, governance, operations, and cost control engineered into the system itself.",
            font_size=17, color=TEXT1, bold=True)
add_textbox(slide, Inches(1.4), Inches(5.7), Inches(10.5), Inches(0.5),
            "This is what Platform Reliability Engineering defines, and what AEGIS operationalizes.",
            font_size=14, color=TEXT3)


# ════════════════════════════════════════════════════════════
# SLIDE 3 — WHAT IS PRE
# ════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)

add_label(slide, Inches(1), Inches(0.6), "DISCIPLINE DEFINITION")
add_textbox(slide, Inches(1), Inches(1.0), Inches(11), Inches(1.0),
            "What is Platform Reliability Engineering?",
            font_size=32, color=TEXT1, bold=True)
add_textbox(slide, Inches(1), Inches(2.0), Inches(8), Inches(0.5),
            "The discipline of ensuring infrastructure platforms remain:",
            font_size=16, color=TEXT2)

# Pillars
pillars = [("Reliable", BRAND_LIGHT), ("Governed", CYAN), ("Cost Efficient", EMERALD),
           ("Secure", AMBER), ("Operable", ROSE), ("Continuously Improving", VIOLET)]
for i, (label, color) in enumerate(pillars):
    col = i % 3
    row = i // 3
    x = Inches(1) + Inches(col * 3.5)
    y = Inches(2.8) + Inches(row * 0.55)
    hex_str = str(color)
    r_val, g_val, b_val = int(hex_str[0:2], 16), int(hex_str[2:4], 16), int(hex_str[4:6], 16)
    r = add_rounded_rect(slide, x, y, Inches(3.2), Inches(0.42),
                         RGBColor(r_val // 6, g_val // 6, b_val // 6))
    add_textbox(slide, x, y + Pt(2), Inches(3.2), Inches(0.38),
                label, font_size=13, color=color, bold=True, alignment=PP_ALIGN.CENTER)

# SRE vs PRE
add_textbox(slide, Inches(1), Inches(4.3), Inches(5), Inches(0.3),
            "WHERE SRE FOCUSES ON", font_size=11, color=TEXT3, bold=True)
add_textbox(slide, Inches(1), Inches(4.6), Inches(5), Inches(0.6),
            "Service reliability \u2014 ensuring individual services meet SLOs and handle failures gracefully.",
            font_size=14, color=TEXT2)

add_textbox(slide, Inches(7), Inches(4.3), Inches(5.5), Inches(0.3),
            "PRE FOCUSES ON", font_size=11, color=BRAND_LIGHT, bold=True)
add_textbox(slide, Inches(7), Inches(4.6), Inches(5.5), Inches(0.8),
            "Platform reliability as a product \u2014 encompassing stability, governance, economics, intelligence, and operations maturity.",
            font_size=14, color=TEXT1, bold=True)

# Bottom pillars
bottom = ["Platform Stability", "Platform Governance", "Platform Economics",
          "Platform Intelligence", "Ops Maturity"]
for i, label in enumerate(bottom):
    x = Inches(1) + Inches(i * 2.35)
    add_rounded_rect(slide, x, Inches(5.8), Inches(2.1), Inches(0.7), SURFACE)
    add_textbox(slide, x, Inches(5.9), Inches(2.1), Inches(0.5),
                label, font_size=12, color=TEXT1, bold=True, alignment=PP_ALIGN.CENTER)


# ════════════════════════════════════════════════════════════
# SLIDE 4 — HOW PRE DIFFERS (TABLE)
# ════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)

add_label(slide, Inches(1), Inches(0.6), "WHY PRE")
add_textbox(slide, Inches(1), Inches(1.0), Inches(11), Inches(1.0),
            "Existing disciplines solve parts of the problem.\nNone own platform operations as a system.",
            font_size=28, color=TEXT1, bold=True)

# Table
rows = 6
cols = 3
tbl_shape = slide.shapes.add_table(rows, cols, Inches(1), Inches(2.5),
                                    Inches(11.3), Inches(3.5))
tbl = tbl_shape.table

headers = ["Discipline", "Primary Focus", "What It Does Not Own"]
data = [
    ("SRE", "Reliability of services", "Cost, governance, platform-level operations"),
    ("Platform Engineering", "Developer experience", "Operational governance, reliability coordination"),
    ("FinOps", "Cloud cost governance", "Reliability, security, operational workflows"),
    ("Cloud Security", "Policy and access control", "Cost, reliability, operational execution"),
    ("PRE", "Unifies all of the above at the platform operations layer", "\u2014"),
]

for i, h in enumerate(headers):
    cell = tbl.cell(0, i)
    cell.text = h
    for p in cell.text_frame.paragraphs:
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = TEXT3
        p.font.name = "Calibri"
    cell.fill.solid()
    cell.fill.fore_color.rgb = SURFACE

for r, (name, focus, gap) in enumerate(data, 1):
    for c, val in enumerate([name, focus, gap]):
        cell = tbl.cell(r, c)
        cell.text = val
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(13)
            p.font.name = "Calibri"
            if r == 5:  # PRE row
                p.font.color.rgb = TEXT1
                p.font.bold = True
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(0x14, 0x18, 0x2E)
            else:
                p.font.color.rgb = TEXT2
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(0x0A, 0x0F, 0x1A)
            if c == 0:
                p.font.bold = True
                p.font.color.rgb = BRAND_LIGHT if r == 5 else TEXT1

add_textbox(slide, Inches(1), Inches(6.2), Inches(11), Inches(0.6),
            "PRE does not replace these disciplines. It is the operating model that aligns them \u2014 ensuring reliability, governance, cost, and security decisions are made together at the platform layer.",
            font_size=14, color=TEXT2)


# ════════════════════════════════════════════════════════════
# SLIDE 5 — FIVE CHALLENGES
# ════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)

add_label(slide, Inches(1), Inches(0.6), "WHY PRE IS EMERGING")
add_textbox(slide, Inches(1), Inches(1.0), Inches(11), Inches(0.8),
            "Five Platform Challenges Driving PRE",
            font_size=32, color=TEXT1, bold=True)

challenges = [
    ("Tool Sprawl", "6+ disconnected tools creating operational fragmentation and knowledge silos.", ROSE),
    ("Operational Toil", "Manual investigations, config fixes, approval workflows consuming bandwidth.", AMBER),
    ("Governance Gaps", "Uncontrolled cloud growth, policy violations, shadow infrastructure.", EMERALD),
    ("Cost Explosion", "Idle resources, overprovisioning, and lack of cost visibility.", AMBER),
    ("No Platform Intelligence", "No health scoring, risk insights, or data-backed operational decisions.", CYAN),
]

for i, (title, desc, color) in enumerate(challenges):
    x = Inches(1) + Inches(i * 2.4)
    add_rounded_rect(slide, x, Inches(2.3), Inches(2.15), Inches(3.5), SURFACE)
    add_accent_bar(slide, x, Inches(2.3), Inches(2.15), color)
    add_textbox(slide, x + Inches(0.15), Inches(2.6), Inches(1.85), Inches(0.5),
                title, font_size=15, color=TEXT1, bold=True)
    add_textbox(slide, x + Inches(0.15), Inches(3.2), Inches(1.85), Inches(1.5),
                desc, font_size=12, color=TEXT2)


# ════════════════════════════════════════════════════════════
# SLIDE 6 — INTRODUCING AEGIS
# ════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)

add_textbox(slide, Inches(1.5), Inches(0.8), Inches(10), Inches(0.4),
            "INTRODUCING AEGIS", font_size=11, color=BRAND_LIGHT, bold=True,
            alignment=PP_ALIGN.CENTER)
add_textbox(slide, Inches(1.5), Inches(1.3), Inches(10), Inches(1.2),
            "AEGIS operationalizes\nPlatform Reliability Engineering",
            font_size=36, color=TEXT1, bold=True, alignment=PP_ALIGN.CENTER)
add_textbox(slide, Inches(2), Inches(2.7), Inches(9), Inches(0.8),
            "AEGIS is the platform reliability control plane that turns PRE from concept into operational reality \u2014 as an independent system first, and an integrated operating layer second.",
            font_size=16, color=TEXT2, alignment=PP_ALIGN.CENTER)

# Statement box
add_rounded_rect(slide, Inches(2), Inches(4.2), Inches(9.3), Inches(2.2), SURFACE, BRAND)
add_accent_bar(slide, Inches(2), Inches(4.2), Inches(9.3), BRAND)
add_textbox(slide, Inches(2.4), Inches(4.5), Inches(8.5), Inches(0.7),
            "Every company that operates complex cloud platforms will eventually need a Platform Reliability Engineering function.",
            font_size=17, color=TEXT1, bold=True, alignment=PP_ALIGN.CENTER)
add_textbox(slide, Inches(2.4), Inches(5.3), Inches(8.5), Inches(0.6),
            "AEGIS brings reliability, governance, operations, and cost into one operational system. It works independently as a platform reliability control plane and integrates with existing tools to provide a unified operating layer.",
            font_size=14, color=TEXT2, alignment=PP_ALIGN.CENTER)


# ════════════════════════════════════════════════════════════
# SLIDE 7 — CONTROL PLANE ARCHITECTURE
# ════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)

add_label(slide, Inches(1), Inches(0.6), "ARCHITECTURE")
add_textbox(slide, Inches(1), Inches(1.0), Inches(11), Inches(1.0),
            "Modern infrastructure has layers.\nOperations did not. Until now.",
            font_size=32, color=TEXT1, bold=True)

layers = [
    ("WORKLOADS", "Applications and services running on your platform", TEXT3, SURFACE, None),
    ("INFRASTRUCTURE", "Cloud accounts, clusters, networks, compute, storage", CYAN, SURFACE, None),
    ("ORCHESTRATION", "Kubernetes, container orchestration, scheduling", EMERALD, SURFACE, None),
    ("OPERATIONS CONTROL PLANE", "AEGIS \u2014 visibility, governance, execution, intelligence", BRAND_LIGHT, RGBColor(0x14, 0x18, 0x2E), BRAND),
]

for i, (label, desc, color, fill, border) in enumerate(layers):
    y = Inches(2.3) + Inches(i * 1.15)
    add_rounded_rect(slide, Inches(2), y, Inches(9), Inches(0.9), fill, border)
    add_textbox(slide, Inches(2.3), y + Inches(0.05), Inches(2.5), Inches(0.4),
                label, font_size=11, color=color, bold=True)
    txt_color = TEXT1 if i == 3 else TEXT2
    add_textbox(slide, Inches(4.8), y + Inches(0.05), Inches(6), Inches(0.8),
                desc, font_size=14, color=txt_color, bold=(i == 3))

add_textbox(slide, Inches(1), Inches(6.2), Inches(10), Inches(0.6),
            "AEGIS is the primary operational system for platform reliability \u2014 able to operate independently through direct platform intelligence, while integrating with your existing stack.",
            font_size=14, color=TEXT2)


# ════════════════════════════════════════════════════════════
# SLIDE 8 — OPERATING MODEL LOOP
# ════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)

add_label(slide, Inches(1), Inches(0.6), "OPERATING MODEL")
add_textbox(slide, Inches(1), Inches(1.0), Inches(11), Inches(0.7),
            "The PRE operational loop", font_size=32, color=TEXT1, bold=True)
add_textbox(slide, Inches(1), Inches(1.8), Inches(8), Inches(0.5),
            "Not features. An operating model. AEGIS enables this continuous loop across your entire platform.",
            font_size=16, color=TEXT2)

steps = [
    ("Discover", "Inventory & baseline"),
    ("Understand", "Signals & context"),
    ("Decide", "Policy evaluation"),
    ("Govern", "Approval & control"),
    ("Execute", "Safe operations"),
    ("Improve", "Intelligence & learning"),
]

for i, (name, desc) in enumerate(steps):
    x = Inches(1) + Inches(i * 2)
    add_rounded_rect(slide, x, Inches(3.0), Inches(1.7), Inches(1.2), SURFACE)
    add_textbox(slide, x, Inches(3.1), Inches(1.7), Inches(0.4),
                name, font_size=15, color=TEXT1, bold=True, alignment=PP_ALIGN.CENTER)
    add_textbox(slide, x, Inches(3.5), Inches(1.7), Inches(0.4),
                desc, font_size=11, color=TEXT3, alignment=PP_ALIGN.CENTER)
    if i < 5:
        add_textbox(slide, x + Inches(1.7), Inches(3.3), Inches(0.3), Inches(0.4),
                    "\u2192", font_size=18, color=TEXT3, alignment=PP_ALIGN.CENTER)

add_textbox(slide, Inches(1), Inches(4.6), Inches(11), Inches(0.4),
            "Every action through this loop produces an immutable audit record.",
            font_size=13, color=TEXT3, alignment=PP_ALIGN.CENTER)

# Back arrow
add_rounded_rect(slide, Inches(1), Inches(5.2), Inches(11.3), Inches(0.6),
                 RGBColor(0x0A, 0x0F, 0x1A))
add_textbox(slide, Inches(1), Inches(5.25), Inches(11.3), Inches(0.5),
            "\u2190  Continuous Loop  \u2192",
            font_size=13, color=BRAND_LIGHT, bold=True, alignment=PP_ALIGN.CENTER)


# ════════════════════════════════════════════════════════════
# SLIDE 9 — FOUR CAPABILITY DOMAINS
# ════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)

add_label(slide, Inches(1), Inches(0.5), "PRE CAPABILITY SYSTEM")
add_textbox(slide, Inches(1), Inches(0.9), Inches(11), Inches(0.7),
            "Four capability domains. One operational system.",
            font_size=28, color=TEXT1, bold=True)

domains = [
    ("01", "FOUNDATION", "Know your platform.", BRAND_LIGHT, BRAND,
     "Resource discovery & cloud inventory\nDrift detection & structural diffing\nIntegration mapping & graph edges\nPlatform baseline visibility\nCanonical normalization & hashing"),
    ("02", "OPERATIONS", "Run your platform reliably.", CYAN, CYAN,
     "Incident coordination & unified triage\nReliability workflows & SLO management\nService intelligence & golden signals\nChange correlation & root cause analysis\nOperational analytics & war rooms"),
    ("03", "CONTROL", "Enforce governance safely.", EMERALD, EMERALD,
     "Policy enforcement & fail-safe decisions\nApproval workflows with SLA deadlines\nRisk visibility & compliance pathways\nExecution guardrails & token validation\nImmutable evidence bundles"),
    ("04", "INTELLIGENCE", "Continuously improve.", AMBER, AMBER,
     "Cost intelligence & anomaly detection\nReliability insights & predictive signals\nPlatform maturity scoring (PRE-100)\nArchitecture risk assessment\nExecutive intelligence & decision support"),
]

for i, (num, label, tagline, lbl_color, accent, features) in enumerate(domains):
    col = i % 2
    row = i // 2
    x = Inches(1) + Inches(col * 5.8)
    y = Inches(1.8) + Inches(row * 2.75)
    add_rounded_rect(slide, x, y, Inches(5.5), Inches(2.5), SURFACE)
    add_accent_bar(slide, x, y, Inches(5.5), accent)

    add_textbox(slide, x + Inches(0.2), y + Inches(0.15), Inches(4), Inches(0.3),
                f"DOMAIN {num} \u2014 {label}", font_size=10, color=lbl_color, bold=True)
    add_textbox(slide, x + Inches(0.2), y + Inches(0.45), Inches(4), Inches(0.35),
                tagline, font_size=16, color=TEXT1, bold=True)
    add_textbox(slide, x + Inches(0.2), y + Inches(0.85), Inches(5), Inches(1.5),
                features, font_size=11, color=TEXT2)


# ════════════════════════════════════════════════════════════
# SLIDE 10 — FIVE PRINCIPLES
# ════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)

add_label(slide, Inches(1), Inches(0.5), "OPERATING PRINCIPLES")
add_textbox(slide, Inches(1), Inches(0.9), Inches(11), Inches(0.7),
            "Five PRE Principles", font_size=32, color=TEXT1, bold=True)

principles = [
    ("01", "Reliability is Engineered",
     "Design failure tolerance, remove single points of failure, define reliability targets.",
     "Reliability cannot depend on hero engineers."),
    ("02", "Governance Must Be Embedded",
     "Governance must exist inside provisioning, change, and deployment workflows.",
     "No change without governance validation."),
    ("03", "Platforms Are Products",
     "Platforms require a roadmap, ownership, SLAs, and experience design.",
     "Treat your platform like a product your teams depend on."),
    ("04", "Automation is Mandatory",
     "Provisioning, scaling, recovery, governance, and incident response must be automated.",
     "If it is repeated, it must be automated."),
    ("05", "Data Drives Decisions",
     "Operational decisions must be data-backed through health scoring and metrics.",
     "Data must guide every platform decision."),
]

for i, (num, title, desc, quote) in enumerate(principles):
    x = Inches(1) + Inches(i * 2.4)
    add_rounded_rect(slide, x, Inches(2.0), Inches(2.15), Inches(4.5), SURFACE)

    add_textbox(slide, x + Inches(0.15), Inches(2.15), Inches(1.85), Inches(0.3),
                f"PRINCIPLE {num}", font_size=10, color=BRAND_LIGHT, bold=True)
    add_textbox(slide, x + Inches(0.15), Inches(2.5), Inches(1.85), Inches(0.6),
                title, font_size=14, color=TEXT1, bold=True)
    add_textbox(slide, x + Inches(0.15), Inches(3.2), Inches(1.85), Inches(1.5),
                desc, font_size=11, color=TEXT2)

    # Quote box
    add_rounded_rect(slide, x + Inches(0.1), Inches(4.7), Inches(1.95), Inches(1.0),
                     RGBColor(0x12, 0x14, 0x2A))
    add_textbox(slide, x + Inches(0.2), Inches(4.8), Inches(1.75), Inches(0.8),
                f'"{quote}"', font_size=10, color=BRAND_LIGHT, bold=True)


# ════════════════════════════════════════════════════════════
# SLIDE 11 — MATURITY MODEL
# ════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)

add_label(slide, Inches(1), Inches(0.6), "PRE-100 MATURITY MODEL")
add_textbox(slide, Inches(1), Inches(1.0), Inches(11), Inches(0.7),
            "Five levels of platform maturity",
            font_size=32, color=TEXT1, bold=True)
add_textbox(slide, Inches(1), Inches(1.8), Inches(8), Inches(0.5),
            "AEGIS moves organizations up this curve \u2014 from reactive firefighting to autonomous platform operations.",
            font_size=16, color=TEXT2)

# Progress bar
bar_y = Inches(3.0)
slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), bar_y,
                       Inches(11.3), Pt(3)).fill.solid()
slide.shapes[-1].fill.fore_color.rgb = RGBColor(0x1A, 0x25, 0x40)
slide.shapes[-1].line.fill.background()

maturity = [
    ("1", "Reactive", "Manual operations, firefighting,\nlimited visibility.", "HIGH RISK", ROSE),
    ("2", "Managed", "Basic monitoring, defined\nprocesses. Human-dependent.", "MODERATE", AMBER),
    ("3", "Standardized", "Defined standards, automation\nintroduced, platform baselines.", "CONSISTENT", CYAN),
    ("4", "Proactive", "Predictive insights, risk detection,\ncost intelligence, reliability scoring.", "PREVENTIVE", EMERALD),
    ("5", "Autonomous", "Systems that continuously detect,\nprioritize, and drive corrective action.", "PRE EVOLUTION", BRAND_LIGHT),
]

for i, (num, name, desc, badge, color) in enumerate(maturity):
    x = Inches(1) + Inches(i * 2.3)
    is_active = num == "5"

    # Dot on line
    dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, x + Inches(0.9),
                                  bar_y - Pt(10), Pt(24), Pt(24))
    dot.fill.solid()
    dot.fill.fore_color.rgb = BRAND if is_active else RGBColor(0x0A, 0x0F, 0x1A)
    dot.line.color.rgb = BRAND if is_active else RGBColor(0x1A, 0x25, 0x40)
    dot.line.width = Pt(2)

    add_textbox(slide, x + Inches(0.85), bar_y - Pt(8), Pt(24), Pt(24),
                num, font_size=10, color=WHITE if is_active else TEXT3,
                bold=True, alignment=PP_ALIGN.CENTER)

    add_textbox(slide, x, Inches(3.5), Inches(2.1), Inches(0.35),
                name, font_size=14, color=TEXT1, bold=True, alignment=PP_ALIGN.CENTER)
    add_textbox(slide, x, Inches(3.85), Inches(2.1), Inches(1.0),
                desc, font_size=11, color=TEXT3, alignment=PP_ALIGN.CENTER)

    # Badge
    hex_s = str(color)
    br, bg_, bb = int(hex_s[0:2], 16), int(hex_s[2:4], 16), int(hex_s[4:6], 16)
    badge_shape = add_rounded_rect(slide, x + Inches(0.35), Inches(5.0),
                                    Inches(1.4), Inches(0.35),
                                    RGBColor(br // 8, bg_ // 8, bb // 8))
    add_textbox(slide, x + Inches(0.35), Inches(5.0), Inches(1.4), Inches(0.35),
                badge, font_size=9, color=color, bold=True, alignment=PP_ALIGN.CENTER)


# ════════════════════════════════════════════════════════════
# SLIDE 12 — PRE STANDARDS
# ════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)

add_label(slide, Inches(1), Inches(0.5), "PRE STANDARDS")
add_textbox(slide, Inches(1), Inches(0.9), Inches(11), Inches(0.7),
            "Platform Standard Areas", font_size=28, color=TEXT1, bold=True)

standards = [
    ("Reliability", BRAND_LIGHT,
     ["SLIs defined for all services", "SLOs enforced with error budgets",
      "Availability & recovery targets", "Failure testing conducted"]),
    ("Governance", EMERALD,
     ["RBAC & least privilege", "Policy enforcement automated",
      "Change approvals required", "Compliance validation continuous"]),
    ("Financial", AMBER,
     ["Cost visibility & allocation", "Anomaly detection",
      "Optimization policies enforced", "Budget enforcement active"]),
    ("Operational Excellence", CYAN,
     ["MTTR tracked and improving", "Change failure rate monitored",
      "Automation coverage measured", "Reliability reviews scheduled"]),
    ("Intelligence", BRAND_LIGHT,
     ["Platform maturity measured", "Reliability posture scored",
      "Risk exposure quantified", "Decision support automated"]),
    ("Key PRE Metrics", BRAND_LIGHT,
     ["Mean Time to Resolve", "Change Failure Rate",
      "Automation Coverage %", "Platform Maturity Score"]),
]

for i, (title, color, items) in enumerate(standards):
    col = i % 3
    row = i // 3
    x = Inches(1) + Inches(col * 3.85)
    y = Inches(1.8) + Inches(row * 2.6)
    add_rounded_rect(slide, x, y, Inches(3.6), Inches(2.35), SURFACE)
    add_textbox(slide, x + Inches(0.2), y + Inches(0.1), Inches(3.2), Inches(0.35),
                title, font_size=14, color=TEXT1, bold=True)
    for j, item in enumerate(items):
        add_textbox(slide, x + Inches(0.2), y + Inches(0.5 + j * 0.4),
                    Inches(3.2), Inches(0.35),
                    f"\u2713  {item}", font_size=11, color=TEXT2)


# ════════════════════════════════════════════════════════════
# SLIDE 13 — EXPECTED OUTCOMES
# ════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)

add_label(slide, Inches(1), Inches(0.6), "EXPECTED OUTCOMES")
add_textbox(slide, Inches(1), Inches(1.0), Inches(11), Inches(0.7),
            "What PRE Delivers", font_size=32, color=TEXT1, bold=True)
add_textbox(slide, Inches(1), Inches(1.8), Inches(8), Inches(0.5),
            "Organizations adopting PRE can expect measurable improvements across four areas.",
            font_size=16, color=TEXT2)

outcomes = [
    ("Reduce Operational Toil", "Automate repetitive workflows, investigations, and manual coordination across teams."),
    ("Improve Platform Uptime", "Proactive risk detection and reliability engineering reduce unplanned outages."),
    ("Reduce Incident MTTR", "Unified context, automated triage, and correlated signals accelerate resolution."),
    ("Improve Cloud Cost Control", "Continuous cost visibility, anomaly detection, and optimization policies."),
]

for i, (title, desc) in enumerate(outcomes):
    x = Inches(1) + Inches(i * 3)
    add_rounded_rect(slide, x, Inches(2.8), Inches(2.75), Inches(2.5), SURFACE)
    add_textbox(slide, x + Inches(0.2), Inches(3.0), Inches(2.35), Inches(0.5),
                title, font_size=15, color=TEXT1, bold=True, alignment=PP_ALIGN.CENTER)
    add_textbox(slide, x + Inches(0.2), Inches(3.6), Inches(2.35), Inches(1.2),
                desc, font_size=12, color=TEXT2, alignment=PP_ALIGN.CENTER)

# Strategic outcome
add_rounded_rect(slide, Inches(2.5), Inches(5.7), Inches(8.3), Inches(1.0),
                 RGBColor(0x14, 0x18, 0x2E), BRAND)
add_textbox(slide, Inches(2.5), Inches(5.75), Inches(8.3), Inches(0.4),
            "Strategic Outcome", font_size=16, color=TEXT1, bold=True,
            alignment=PP_ALIGN.CENTER)
add_textbox(slide, Inches(2.5), Inches(6.15), Inches(8.3), Inches(0.4),
            "Your platform becomes safer, faster, and cheaper to operate.",
            font_size=14, color=BRAND_LIGHT, bold=True, alignment=PP_ALIGN.CENTER)


# ════════════════════════════════════════════════════════════
# SLIDE 14 — AEGIS STANDALONE
# ════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)

add_label(slide, Inches(1), Inches(0.6), "CLARITY")
add_textbox(slide, Inches(1), Inches(1.0), Inches(11), Inches(1.0),
            "AEGIS works independently \u2014\nand gets stronger with integrations",
            font_size=28, color=TEXT1, bold=True)

add_textbox(slide, Inches(1), Inches(2.3), Inches(5.5), Inches(2),
            "AEGIS is not dependent on external tools to deliver value. It operates as a standalone platform reliability control plane using direct platform intelligence, governance logic, workflows, and operational decisioning.\n\nWhen connected to your monitoring, security, cost, and incident stack, it becomes the unified operating layer across your environment.",
            font_size=14, color=TEXT2)

# Right side box
add_rounded_rect(slide, Inches(7), Inches(2.3), Inches(5.3), Inches(4.5), SURFACE)
add_textbox(slide, Inches(7.3), Inches(2.4), Inches(4.7), Inches(0.35),
            "AEGIS STANDALONE CORE", font_size=10, color=TEXT3, bold=True)

core_items = [
    "Platform discovery & baseline visibility",
    "Governance workflows & policy decisions",
    "Operational coordination & audit trails",
    "Cost, reliability, and maturity intelligence",
]
for i, item in enumerate(core_items):
    y = Inches(2.9) + Inches(i * 0.55)
    add_rounded_rect(slide, Inches(7.2), y, Inches(4.9), Inches(0.45),
                     RGBColor(0x0A, 0x0F, 0x1A))
    add_textbox(slide, Inches(7.3), y + Pt(2), Inches(4.7), Inches(0.4),
                f"\u2713  {item}", font_size=12, color=TEXT2)

add_rounded_rect(slide, Inches(7.2), Inches(5.2), Inches(4.9), Inches(0.5),
                 RGBColor(0x14, 0x18, 0x2E), BRAND)
add_textbox(slide, Inches(7.3), Inches(5.25), Inches(4.7), Inches(0.4),
            "Existing tools become inputs and outputs to the AEGIS operating layer",
            font_size=11, color=TEXT1, bold=True)


# ════════════════════════════════════════════════════════════
# SLIDE 15 — DESIGN PARTNER PROGRAM
# ════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)

# Big card
add_rounded_rect(slide, Inches(1), Inches(1), Inches(11.3), Inches(5.5),
                 RGBColor(0x0D, 0x12, 0x22), BRAND)
add_accent_bar(slide, Inches(1), Inches(1), Inches(11.3), BRAND)

add_textbox(slide, Inches(1.5), Inches(1.5), Inches(5), Inches(0.3),
            "DESIGN PARTNER PROGRAM", font_size=11, color=BRAND_LIGHT, bold=True)
add_textbox(slide, Inches(1.5), Inches(2.0), Inches(5.5), Inches(1.2),
            "Shape the future of\nPlatform Reliability Engineering",
            font_size=28, color=TEXT1, bold=True)
add_textbox(slide, Inches(1.5), Inches(3.3), Inches(5), Inches(1),
            "We are working with a limited number of platform teams to shape AEGIS. If you are building serious platform capabilities, we want to work with you.",
            font_size=15, color=TEXT2)

# Benefits
benefits = ["Early product access", "Direct roadmap influence",
            "Architecture collaboration", "Founder access", "Preferred pricing"]
for i, b in enumerate(benefits):
    y = Inches(1.8) + Inches(i * 0.65)
    add_rounded_rect(slide, Inches(7.5), y, Inches(4.3), Inches(0.5),
                     RGBColor(0x0A, 0x0F, 0x1A))
    add_textbox(slide, Inches(7.6), y + Pt(2), Inches(4.1), Inches(0.45),
                f"\u2713  {b}", font_size=14, color=TEXT2)


# ════════════════════════════════════════════════════════════
# SLIDE 16 — CLOSING CTA
# ════════════════════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide)

add_textbox(slide, Inches(1.5), Inches(1.0), Inches(10), Inches(0.4),
            "BRAHMORA TECHNOLOGIES", font_size=13, color=BRAND_LIGHT, bold=True,
            alignment=PP_ALIGN.CENTER)

add_textbox(slide, Inches(1.5), Inches(2.0), Inches(10), Inches(1.5),
            "Platform reliability is becoming\na discipline.",
            font_size=42, color=TEXT1, bold=True, alignment=PP_ALIGN.CENTER)

add_textbox(slide, Inches(2.5), Inches(3.8), Inches(8), Inches(0.6),
            "PRE defines it. AEGIS enables it.\nJoin the companies shaping this future.",
            font_size=18, color=TEXT2, alignment=PP_ALIGN.CENTER)

# CTA button shape
btn = add_rounded_rect(slide, Inches(4.5), Inches(4.8), Inches(4.3), Inches(0.7),
                       BRAND)
add_textbox(slide, Inches(4.5), Inches(4.85), Inches(4.3), Inches(0.6),
            "Become a Design Partner  \u2192", font_size=16, color=WHITE,
            bold=True, alignment=PP_ALIGN.CENTER)

add_textbox(slide, Inches(1.5), Inches(6.2), Inches(10), Inches(0.4),
            "brahmora.co.uk", font_size=14, color=TEXT3, alignment=PP_ALIGN.CENTER)


# ════════════════════════════════════════════════════════════
# SAVE
# ════════════════════════════════════════════════════════════
output_path = "/Users/lsakthivel/brahmora/brahmora-wp/Brahmora-AEGIS-PRE.pptx"
prs.save(output_path)
print(f"Saved: {output_path}")
